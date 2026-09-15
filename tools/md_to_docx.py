#!/usr/bin/env python3
"""
Render the project's Markdown documents as Word .docx files.

Markdown under docs/ is the source of truth; this script produces the printable
submission copies in docs/docx/. It handles the subset of Markdown the project
documents use: headings, paragraphs, pipe tables, fenced code blocks, bullet and
numbered lists, horizontal rules, and inline bold, italic, code and links.

A page break is inserted after the Team Members table so that the document
control block and the team names and IDs occupy the first page of every document.

Usage:
    pip install python-docx
    python3 tools/md_to_docx.py                 # all documents
    python3 tools/md_to_docx.py docs/test_plan.md

Author: Henok Zemedkun (ATE/8552/16) - Person 2, Test Architect
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / 'docs'
OUT_DIR = DOCS / 'docx'

DEFAULT_DOCUMENTS = [
    'test_plan.md',
    'test_design.md',
    'test_summary_report.md',
    'foundations_reflection.md',
]

HEADING_LEVELS = {2: 1, 3: 2, 4: 3, 5: 4, 6: 4}
INLINE = re.compile(
    r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\n]+\*(?![\w*]))'
)


# --------------------------------------------------------------------------
# Inline formatting
# --------------------------------------------------------------------------

def add_runs(paragraph, text, bold=False, size=None):
    """Append `text` to `paragraph`, honouring inline Markdown formatting."""
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif token.startswith('[') and '](' in token:
            label = token[1:token.index('](')]
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x1A, 0x4F, 0xA0)
            run.underline = True
        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)


# --------------------------------------------------------------------------
# Block elements
# --------------------------------------------------------------------------

def split_row(line):
    """Split a Markdown table row into its cells."""
    return [cell.strip() for cell in line.strip().strip('|').split('|')]


def is_separator(line):
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', line.strip()))


def add_table(document, rows):
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=len(rows), cols=len(header))
    table.style = 'Table Grid'
    table.autofit = True

    for index, cell_text in enumerate(header):
        cell = table.cell(0, index)
        cell.text = ''
        add_runs(cell.paragraphs[0], cell_text, bold=True, size=9)

    for row_index, row in enumerate(body, start=1):
        for column_index in range(len(header)):
            cell = table.cell(row_index, column_index)
            cell.text = ''
            value = row[column_index] if column_index < len(row) else ''
            add_runs(cell.paragraphs[0], value, size=9)

    document.add_paragraph()
    return table


def add_code_block(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run('\n'.join(lines))
    run.font.name = 'Consolas'
    run.font.size = Pt(8)


def convert(md_path, out_path):
    lines = md_path.read_text(encoding='utf-8').splitlines()
    document = Document()

    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)

    title_done = False
    team_table_seen = False
    page_break_pending = False
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith('```'):
            fence = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                fence.append(lines[index])
                index += 1
            add_code_block(document, fence)
            index += 1
            continue

        # table
        if stripped.startswith('|') and index + 1 < len(lines) \
                and is_separator(lines[index + 1]):
            rows = [split_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(split_row(lines[index]))
                index += 1
            table = add_table(document, rows)
            header_text = ' '.join(rows[0]).lower()
            if 'student id' in header_text:
                team_table_seen = True
                page_break_pending = True
            del table
            continue

        # heading
        heading = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if heading:
            level, text = len(heading.group(1)), heading.group(2)
            if level == 1 and not title_done:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(text)
                run.bold = True
                run.font.size = Pt(20)
                document.add_paragraph()
                title_done = True
            else:
                if page_break_pending and team_table_seen:
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    page_break_pending = False
                document.add_heading(text, level=HEADING_LEVELS.get(level, 4))
            index += 1
            continue

        # horizontal rule
        if stripped in ('---', '***', '___'):
            index += 1
            continue

        # bullet list
        bullet = re.match(r'^\s*[-*]\s+(.*)', line)
        if bullet:
            paragraph = document.add_paragraph(style='List Bullet')
            add_runs(paragraph, bullet.group(1))
            index += 1
            continue

        # numbered list
        numbered = re.match(r'^\s*\d+\.\s+(.*)', line)
        if numbered:
            paragraph = document.add_paragraph(style='List Number')
            add_runs(paragraph, numbered.group(1))
            index += 1
            continue

        # blank line
        if not stripped:
            index += 1
            continue

        # paragraph: join the following non-blank, non-block lines
        buffer = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if not nxt or nxt.startswith(('|', '#', '```', '---')) \
                    or re.match(r'^\s*([-*]|\d+\.)\s+', lines[index]):
                break
            buffer.append(nxt)
            index += 1
        add_runs(document.add_paragraph(), ' '.join(buffer))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    return out_path


def main():
    targets = sys.argv[1:] or [str(DOCS / name) for name in DEFAULT_DOCUMENTS]
    for target in targets:
        md_path = Path(target)
        if not md_path.is_absolute():
            md_path = ROOT / md_path
        out_path = OUT_DIR / (md_path.stem + '.docx')
        convert(md_path, out_path)
        size_kb = out_path.stat().st_size / 1024
        print(f'{md_path.relative_to(ROOT)} -> '
              f'{out_path.relative_to(ROOT)} ({size_kb:.0f} KB)')


if __name__ == '__main__':
    main()
